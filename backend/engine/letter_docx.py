"""The clarification request as a Word letter (python-docx).

Same questions, same quotes and same letter texts as the plain-text draft (engine.recommend.clarification_request);
only the layout differs. Nothing is added to the questions: the requirement codes come from the question sources,
and a question without a code is written without one. The translator is passed in, so this module does not depend
on the web layer.
"""

from __future__ import annotations

import io
from collections.abc import Callable, Iterable
from datetime import date
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from engine.recommend import _LETTER_HEAD, QuestionView

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_INDENT = Cm(0.75)


def _codes(question: QuestionView) -> list[str]:
    """Requirement codes of the question's sources, in order, each once."""
    codes: list[str] = []
    for src in question.sources:
        code = str(src.get("code") or "").strip()
        if code and code not in codes:
            codes.append(code)
    return codes


def _where(src: dict[str, Any], file_label: str, page_label: str) -> str:
    parts = []
    if src.get("file"):
        parts.append(f"{file_label} {src['file']}")
    if src.get("page") is not None:
        parts.append(f"{page_label} {src['page']}")
    return ", ".join(parts)


def clarification_docx(
    questions: Iterable[QuestionView],
    *,
    lang: str,
    tr: Callable[..., str],
    project_name: str,
    customer: str | None,
    lot_number: str | None,
    today: date,
) -> bytes:
    """The clarification letter as .docx bytes. `tr` is the UI translator for `lang` (keys clarification.*)."""
    head_lang = lang if lang in _LETTER_HEAD else "ru"
    title, intro, quote_label, page_label, file_label = _LETTER_HEAD[head_lang]

    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = "Tender-Sverka"

    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)  # A4
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2))

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.add_run(title).bold = True
    heading.runs[0].font.size = Pt(14)

    addressee = doc.add_paragraph()
    addressee.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    addressee.add_run(tr("clarification.to", name=customer) if customer else tr("clarification.to_customer"))

    doc.add_paragraph(tr("clarification.project", name=project_name))
    if lot_number:
        doc.add_paragraph(tr("clarification.lot", lot=lot_number))
    doc.add_paragraph(tr("clarification.date", date=today.strftime("%d.%m.%Y")))

    doc.add_paragraph(intro)

    items = list(questions)
    if not items:
        doc.add_paragraph(tr("clarification.no_questions"))
    for number, question in enumerate(items, start=1):
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = _INDENT
        para.paragraph_format.first_line_indent = -_INDENT
        para.paragraph_format.space_before = Pt(6)
        para.add_run(f"{number}.\t")
        codes = _codes(question)
        if codes:
            para.add_run(", ".join(codes)).bold = True
            para.add_run(" — ")
        para.add_run(question.text)

        for src in question.sources:
            quote = str(src.get("quote") or "").strip()
            if not quote:
                continue
            where = _where(src, file_label, page_label)
            source = doc.add_paragraph()
            source.paragraph_format.left_indent = _INDENT
            source.add_run(f"{quote_label}" + (f" ({where})" if where else "") + ": ")
            source.add_run(f"«{quote}»").italic = True

    closing = doc.add_paragraph(tr("clarification.closing"))
    closing.paragraph_format.space_before = Pt(18)
    signature = doc.add_paragraph(tr("clarification.signature"))
    signature.paragraph_format.space_before = Pt(24)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
