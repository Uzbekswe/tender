"""Customer-side API (docs/customer-side-decisions.md §4-7) with synthetic data.

Fake finished runs are written straight into the temp DB (like tests/unit/test_results_api.py), then
scored with recompute_run. Check-all tests stub runner.enqueue, so no pipeline stage and no model
call ever runs. Expected numbers are worked out in the comments of the seed.
"""

from __future__ import annotations

import io
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlmodel import select
from sqlalchemy import text

from app import runner as runner_module
from app.db import get_engine, init_db, session_scope
from app.models import (
    AuditLog,
    ChecklistVersion,
    Criterion,
    Document,
    Evidence,
    EvaluationScheme,
    MatchResult,
    Package,
    PackageRole,
    Page,
    PageSource,
    Project,
    Requirement,
    ReqType,
    Run,
    RunKind,
    RunStatus,
    Scope,
    utcnow,
)
from engine.score_stages import recompute_run

API = "/api/v1"
ADVISORY_EN = "Advisory: the evaluation committee decides"
QUOTE = "SYNTHETIC QUOTE ZEBRA 42 is the registration certificate"

BIDDERS = {
    "A": "Bidder Alpha",
    "B": "Bidder Beta",
    "C": "Bidder Gamma",
    "D": "Bidder Delta",   # never ran
    "E": "Bidder Epsilon",
}
STATUSES = {
    # code: (status, reason codes) per bidder
    "A": {"T-001": ("meets", ["VALUE_OK"]), "U-001": ("meets", []), "U-002": ("meets", []),
          "U-003": ("meets", []), "U-004": ("needs_review", [])},
    "B": {"T-001": ("meets", []), "U-001": ("needs_review", ["UNREADABLE_IN_SCOPE"]),
          "U-002": ("partial", ["PARTIAL_COVERAGE"]), "U-003": ("meets", []), "U-004": ("needs_review", [])},
    "C": {"T-001": ("meets", []), "U-001": ("does_not_meet", ["ABSENT"]), "U-002": ("meets", []),
          "U-003": ("not_confirmed", ["CLAIM_ONLY"]), "U-004": ("needs_review", [])},
    "E": {"T-001": ("meets", []), "U-001": ("needs_review", ["UNREADABLE_IN_SCOPE"]),
          "U-002": ("partial", ["PARTIAL_COVERAGE"]), "U-003": ("needs_review", ["UNREADABLE_IN_SCOPE"]),
          "U-004": ("needs_review", [])},
}
LEVELS = {4: [4, 2, 0], 2: [2, 0]}


# ---------- seeding ----------

def _bare_customer(*, approved: bool, lang: str = "en") -> dict[str, Any]:
    """A customer project with a tender, an approved (or draft) checklist and a bidder without files."""
    with session_scope() as s:
        project = Project(name="Lot bare", report_lang=lang, mode="customer")
        s.add(project)
        s.flush()
        tender = Package(project_id=project.id, role=PackageRole.tender, label="Tender")
        s.add(tender)
        s.flush()
        tender_run = Run(package_id=tender.id, kind=RunKind.tender, status=RunStatus.waiting_approval)
        s.add(tender_run)
        s.flush()
        version = ChecklistVersion(
            project_id=project.id, number=1, status="approved" if approved else "draft",
            source_run_id=tender_run.id, approved_at=utcnow() if approved else None,
        )
        s.add(version)
        s.flush()
        return {"project": project.id, "version": version.id, "tender": tender.id}


def _seed(*, scheme: bool = True, lang: str = "en") -> dict[str, Any]:
    """One approved checklist, five bidders (A to E), D has no run. Expected results:

    A: meets all (mandatory 2 of 2). Points 6 of 6. Compliance 100 %. Verdict ready. Rank 1.
    B: open points (U-001 needs review). Points 4. Compliance 75 %, range 54.5 to 81.8. Needs revision.
    E: open points. Points 2 to 4 (pass mark 4). Compliance 71.4 %, range 45.5 to 81.8. Needs revision.
       B and E: points ranges touch (E reaches 4, B has 4), so both are flagged "order not certain".
    C: U-001 absent and mandatory. Compliance 66.4 %. Subject to rejection, no rank.
    """
    with session_scope() as s:
        project = Project(name="Lot 9", report_lang=lang, mode="customer")
        s.add(project)
        s.flush()
        tender = Package(project_id=project.id, role=PackageRole.tender, label="Tender")
        s.add(tender)
        s.flush()
        tdoc = Document(package_id=tender.id, file_name="tender.pdf", stored_path="unused", sha256="d" * 64)
        s.add(tdoc)
        s.flush()
        tpage = Page(document_id=tdoc.id, number=5, text="Rated power not less than 2.5 kW.", source=PageSource.text_layer)
        s.add(tpage)
        s.flush()
        tender_run = Run(package_id=tender.id, kind=RunKind.tender, status=RunStatus.waiting_approval)
        s.add(tender_run)
        s.flush()
        version = ChecklistVersion(project_id=project.id, number=1, status="approved",
                                   source_run_id=tender_run.id, approved_at=utcnow())
        s.add(version)
        s.flush()

        exp = warr = None
        if scheme:
            sch = EvaluationScheme(checklist_version_id=version.id, stated_max=6, pass_threshold=4,
                                   platform_scored=[{"name": "price", "points": 25}], source_page_id=tpage.id)
            s.add(sch)
            s.flush()
            exp = Criterion(scheme_id=sch.id, number=1, name="Experience", max_points=4,
                            levels=[{"points": p, "condition_text": "", "rule": None} for p in LEVELS[4]])
            warr = Criterion(scheme_id=sch.id, number=2, name="Warranty", max_points=2,
                             levels=[{"points": p, "condition_text": "", "rule": None} for p in LEVELS[2]])
            s.add_all([exp, warr])
            s.flush()

        reqs: dict[str, Requirement] = {}
        specs = [
            ("T-001", Scope.goods, "Technical", "Rated power", ReqType.mandatory, dict(operator=">=", value="2.5", unit="kW")),
            ("U-001", Scope.bidder, "Documents", "Registration", ReqType.mandatory,
             dict(confirmation_method="Registration certificate")),
            ("U-002", Scope.bidder, "Qualification", "Experience", ReqType.scored,
             dict(criterion_id=exp.id if exp else None)),
            ("U-003", Scope.bidder, "Commercial", "Warranty", ReqType.desirable,
             dict(operator=">=", value="24", unit="months", criterion_id=warr.id if warr else None)),
            ("U-004", Scope.bidder, "Documents", "Digital signature", ReqType.mandatory, dict(verification="platform")),
            ("U-005", Scope.bidder, "Documents", "Excluded item", ReqType.desirable, dict(excluded=True)),
        ]
        for code, scope, category, parameter, req_type, extra in specs:
            req = Requirement(
                checklist_version_id=version.id, code=code, scope=scope, category=category,
                parameter=parameter, req_type=req_type, source_page_id=tpage.id,
                source_quote=f"{parameter} requirement", **extra,
            )
            s.add(req)
            s.flush()
            reqs[code] = req

        packages: dict[str, Package] = {}
        documents: dict[str, Document] = {}
        for key, name in BIDDERS.items():
            pkg = Package(project_id=project.id, role=PackageRole.bid, label=name, bidder_name=name,
                          bidder_stir="123456789" if key == "A" else None)
            s.add(pkg)
            s.flush()
            doc = Document(package_id=pkg.id, file_name=f"{key}.pdf", stored_path="unused", sha256=key * 64)
            s.add(doc)
            s.flush()
            packages[key] = pkg
            documents[key] = doc

        run_ids: dict[str, int] = {}
        for key in ("A", "B", "C", "E"):
            run = Run(package_id=packages[key].id, kind=RunKind.bid, status=RunStatus.done,
                      checklist_version_id=version.id, model_config_used={"match": "m1"},
                      prompt_versions={"match": "match@v1"}, reproducible=True)
            s.add(run)
            s.flush()
            run_ids[key] = run.id
            bid_page = Page(document_id=documents[key].id, number=1, text=QUOTE, source=PageSource.text_layer)
            s.add(bid_page)
            s.flush()
            for code, (status, reasons) in STATUSES[key].items():
                match = MatchResult(run_id=run.id, requirement_id=reqs[code].id, status=status,
                                    reason_codes=reasons, confidence=90, method="llm")
                s.add(match)
                s.flush()
                if key == "A" and code == "U-001":
                    s.add(Evidence(match_result_id=match.id, page_id=bid_page.id, quote=QUOTE,
                                   quote_verified=True, verify_method="exact", subject_is_bidder=True))
        ids = {
            "project": project.id, "tender": tender.id, "version": version.id,
            "packages": {k: p.id for k, p in packages.items()}, "runs": run_ids,
        }
    for run_id in run_ids.values():
        recompute_run(run_id)
    return ids


@pytest.fixture
def seeded(client: TestClient) -> dict[str, Any]:
    return _seed()


def _pkg(ids: dict[str, Any], key: str) -> int:
    return ids["packages"][key]


def _comparison(client: TestClient, project_id: int, query: str = "") -> dict[str, Any]:
    resp = client.get(f"{API}/projects/{project_id}/comparison{query}")
    assert resp.status_code == 200, resp.text
    return resp.json()


def _bidder(body: dict[str, Any], package_id: int) -> dict[str, Any]:
    return next(b for b in body["bidders"] if b["package_id"] == package_id)


# ---------- project mode ----------

def test_mode_defaults_to_supplier_and_can_be_set(client: TestClient):
    default = client.post(f"{API}/projects", json={"name": "Old style"})
    assert default.status_code == 201 and default.json()["mode"] == "supplier"

    customer = client.post(f"{API}/projects", json={"name": "Tender", "mode": "customer"})
    assert customer.status_code == 201 and customer.json()["mode"] == "customer"
    assert client.get(f"{API}/projects/{customer.json()['id']}").json()["mode"] == "customer"

    assert client.post(f"{API}/projects", json={"name": "Bad", "mode": "other"}).status_code == 422


def test_mode_column_is_added_to_an_existing_database(data_dir):
    init_db()
    with get_engine().begin() as conn:
        conn.execute(text("ALTER TABLE project DROP COLUMN mode"))
        conn.execute(text(
            "INSERT INTO project (company_id, name, report_lang, settings, created_at) "
            "VALUES (1, 'Old lot', 'ru', '{}', '2026-01-01 00:00:00')"
        ))
    init_db()  # the MVP migration adds the column back with its default
    with get_engine().connect() as conn:
        columns = [row[1] for row in conn.execute(text("PRAGMA table_info(project)"))]
        assert "mode" in columns
        assert conn.execute(text("SELECT mode FROM project WHERE name = 'Old lot'")).scalar() == "supplier"


# ---------- bidders ----------

def test_bidder_create_list_and_stir_validation(client: TestClient):
    project = client.post(f"{API}/projects", json={"name": "Lot", "mode": "customer"}).json()
    url = f"{API}/projects/{project['id']}/bidders"

    created = client.post(url, json={"name": "ООО Тест", "stir": "123456789"})
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["role"] == "bid" and body["bidder_name"] == "ООО Тест"
    assert body["bidder_stir"] == "123456789" and body["label"] == "ООО Тест"

    assert client.post(url, json={"name": "Short", "stir": "12345678"}).status_code == 422
    assert client.post(url, json={"name": "Letters", "stir": "1234567a9"}).status_code == 422
    assert client.post(url, json={"name": "   "}).status_code == 422
    blank = client.post(url, json={"name": "No STIR", "stir": ""})
    assert blank.status_code == 201 and blank.json()["bidder_stir"] is None
    padded = client.post(url, json={"name": "Padded", "stir": " 987654321 "})
    assert padded.json()["bidder_stir"] == "987654321"

    names = [b["bidder_name"] for b in client.get(url).json()]
    assert names == ["ООО Тест", "No STIR", "Padded"]
    assert client.get(f"{API}/projects/999/bidders").status_code == 404


def test_bidder_patch_and_delete(client: TestClient, data_dir):
    project = client.post(f"{API}/projects", json={"name": "Lot", "mode": "customer"}).json()
    tender = client.post(f"{API}/projects/{project['id']}/packages", json={"role": "tender", "label": "Tender"}).json()
    url = f"{API}/projects/{project['id']}/bidders"
    bidder = client.post(url, json={"name": "First", "stir": "111111111"}).json()
    bid_url = f"{url}/{bidder['id']}"

    patched = client.patch(bid_url, json={"name": "Renamed"})
    assert patched.status_code == 200 and patched.json()["label"] == "Renamed"
    assert patched.json()["bidder_stir"] == "111111111"  # untouched

    assert client.patch(bid_url, json={"stir": None}).json()["bidder_stir"] is None
    assert client.patch(bid_url, json={"stir": "12345678"}).status_code == 422
    assert client.patch(bid_url, json={"name": None}).status_code == 422
    assert client.patch(f"{url}/{tender['id']}", json={"name": "Nope"}).status_code == 404  # not a bidder

    assert client.delete(f"{url}/{tender['id']}").status_code == 404
    assert client.delete(bid_url).status_code == 204
    assert client.get(url).json() == []


def test_bidder_with_a_finished_check_cannot_be_deleted(seeded, client: TestClient):
    url = f"{API}/projects/{seeded['project']}/bidders/{_pkg(seeded, 'A')}"
    assert client.delete(url).status_code == 409
    with session_scope() as s:
        assert s.get(Package, _pkg(seeded, "A")) is not None  # nothing was removed


def test_bidder_with_a_check_in_progress_cannot_be_deleted(seeded, client: TestClient):
    with session_scope() as s:
        pkg = Package(project_id=seeded["project"], role=PackageRole.bid, label="Queued", bidder_name="Queued")
        s.add(pkg)
        s.flush()
        s.add(Run(package_id=pkg.id, kind=RunKind.bid, status=RunStatus.queued,
                  checklist_version_id=seeded["version"]))
        pkg_id = pkg.id
    assert client.delete(f"{API}/projects/{seeded['project']}/bidders/{pkg_id}").status_code == 409


def test_bidder_with_files_but_no_check_is_deleted_with_its_uploads(client: TestClient, seeded):
    url = f"{API}/projects/{seeded['project']}/bidders/{_pkg(seeded, 'D')}"
    assert client.delete(url).status_code == 204
    with session_scope() as s:
        assert s.get(Package, _pkg(seeded, "D")) is None
        assert s.exec(select(Document).where(Document.package_id == _pkg(seeded, "D"))).all() == []
        assert len(s.exec(select(AuditLog).where(AuditLog.action == "bidder.delete")).all()) == 1


# ---------- check all ----------

def test_check_all_needs_an_approved_checklist(client: TestClient):
    ids = _bare_customer(approved=False)
    assert client.post(f"{API}/projects/{ids['project']}/check-all").status_code == 409


def test_check_all_starts_one_run_per_bidder_with_files_on_one_version(client: TestClient, monkeypatch, data_dir):
    queued: list[int] = []
    monkeypatch.setattr(runner_module.runner, "enqueue", lambda run_id: queued.append(run_id))
    ids = _bare_customer(approved=True)
    url = f"{API}/projects/{ids['project']}/bidders"
    with_files = [client.post(url, json={"name": f"Bidder {n}"}).json()["id"] for n in ("one", "three")]
    no_files = client.post(url, json={"name": "Two"}).json()["id"]
    for pkg_id in with_files:
        upload = client.post(f"{API}/packages/{pkg_id}/files",
                             files=[("files", ("offer.pdf", b"%PDF-1.4 synthetic offer", "application/pdf"))])
        assert upload.status_code == 201, upload.text

    resp = client.post(f"{API}/projects/{ids['project']}/check-all")
    assert resp.status_code == 202, resp.text
    body = resp.json()
    assert body["checklist_version_id"] == ids["version"]
    assert [r["package_id"] for r in body["runs"]] == with_files
    assert [s["package_id"] for s in body["skipped"]] == [no_files]
    assert body["skipped"][0]["reason"] == "no files uploaded"
    assert queued == [r["run_id"] for r in body["runs"]]
    for run in body["runs"]:
        row = client.get(f"{API}/runs/{run['run_id']}").json()
        assert row["kind"] == "bid" and row["checklist_version_id"] == ids["version"]


def test_check_all_skips_a_bidder_already_in_progress_and_reruns_on_a_new_version(client: TestClient, monkeypatch, data_dir):
    queued: list[int] = []
    monkeypatch.setattr(runner_module.runner, "enqueue", lambda run_id: queued.append(run_id))
    ids = _bare_customer(approved=True)
    pkg = client.post(f"{API}/projects/{ids['project']}/bidders", json={"name": "Only"}).json()["id"]
    client.post(f"{API}/packages/{pkg}/files", files=[("files", ("a.pdf", b"%PDF synthetic", "application/pdf"))])
    url = f"{API}/projects/{ids['project']}/check-all"

    first = client.post(url).json()
    assert len(first["runs"]) == 1
    again = client.post(url).json()  # the first run is still queued
    assert again["runs"] == []
    assert again["skipped"][0]["reason"] == "a check of this bidder is already running"

    with session_scope() as s:  # a later approved checklist version re-runs every bidder
        s.add(ChecklistVersion(project_id=ids["project"], number=2, status="approved", approved_at=utcnow()))
    later = client.post(url).json()
    assert later["checklist_version_number"] == 2
    assert len(later["runs"]) == 1
    assert len(queued) == 2


def test_check_all_with_no_bidders_is_refused(client: TestClient):
    ids = _bare_customer(approved=True)
    assert client.post(f"{API}/projects/{ids['project']}/check-all").status_code == 409


# ---------- comparison: payload ----------

def test_comparison_without_an_approved_checklist_is_409(client: TestClient):
    ids = _bare_customer(approved=False)
    assert client.get(f"{API}/projects/{ids['project']}/comparison").status_code == 409


def test_comparison_top_level_and_ranking(client: TestClient, seeded):
    body = _comparison(client, seeded["project"])
    assert body["advisory_label"] == ADVISORY_EN
    assert body["checklist_version"]["id"] == seeded["version"]
    assert body["checklist_version"]["number"] == 1

    ranked = [(b["name"], b["rank"]) for b in body["bidders"]]
    assert ranked == [
        (BIDDERS["A"], 1),
        (BIDDERS["B"], 2),
        (BIDDERS["E"], 3),
        (BIDDERS["C"], None),  # subject to rejection: last, no rank
        (BIDDERS["D"], None),  # never ran: last
    ]


def test_comparison_bidder_numbers_match_the_hand_computed_values(client: TestClient, seeded):
    body = _comparison(client, seeded["project"])
    a = _bidder(body, _pkg(seeded, "A"))
    assert a["verdict"] == "ready" and a["verdict_committee_text"] == "Meets all mandatory requirements (advisory)"
    assert a["compliance"] == {"min": 100.0, "max": 100.0, "point": 100.0}
    assert a["points"]["min"] == 6.0 and a["points"]["max"] == 6.0
    assert a["points"]["pass_mark"] == 4.0 and a["points"]["verdict"] == "pass"
    assert a["points"]["note"] == "non-price points only"  # the scheme has platform-scored parts
    assert (a["mandatory_met"], a["mandatory_total"]) == (2, 2)
    assert a["counts"] == {"meets": 4, "partial": 0, "not_confirmed": 0, "not_met": 0, "needs_review": 0}
    assert a["order_not_certain"] is False and a["reproducible"] is True

    b = _bidder(body, _pkg(seeded, "B"))
    assert b["verdict"] == "needs_revision"
    assert b["compliance"]["point"] == 75.0
    assert b["compliance"]["min"] == 54.5 and b["compliance"]["max"] == 81.8
    assert b["points"]["min"] == 4.0 and b["points"]["max"] == 4.0
    assert b["counts"] == {"meets": 2, "partial": 1, "not_confirmed": 0, "not_met": 0, "needs_review": 1}
    assert b["order_not_certain"] is True

    e = _bidder(body, _pkg(seeded, "E"))
    assert e["points"]["min"] == 2.0 and e["points"]["max"] == 4.0
    assert e["points"]["verdict"] == "uncertain"
    assert e["order_not_certain"] is True

    c = _bidder(body, _pkg(seeded, "C"))
    assert c["verdict"] == "rejection_risk"
    assert c["verdict_committee_text"] == "Subject to rejection by the committee"
    assert c["compliance"]["point"] == 66.4
    assert c["rank"] is None and c["order_not_certain"] is False


def test_deciding_items_carry_quotes_and_committee_sentences(client: TestClient, seeded):
    body = _comparison(client, seeded["project"])
    c = _bidder(body, _pkg(seeded, "C"))
    assert [d["code"] for d in c["deciding_items"]] == ["U-001"]
    decided = c["deciding_items"][0]
    assert decided["status"] == "does_not_meet"
    assert decided["text"] == (
        "Bidder did not provide a document confirming Registration. "
        "Does not meet mandatory requirement, subject to rejection by the committee."
    )
    assert decided["quote"] == "" and decided["page"] is None  # the bid has no evidence for it

    a = _bidder(body, _pkg(seeded, "A"))
    u001 = next(d for d in a["deciding_items"] if d["code"] == "U-001")
    assert u001["quote"] == QUOTE and u001["page"] == 1 and u001["doc"] == "A.pdf"


def test_rows_are_grouped_and_marked(client: TestClient, seeded):
    body = _comparison(client, seeded["project"])
    codes = [r["code"] for r in body["rows"]]
    # Grouped by scope, then category (Documents first: it has the first code), U-005 excluded.
    assert codes == ["T-001", "U-001", "U-004", "U-002", "U-003"]
    rows = {r["code"]: r for r in body["rows"]}
    assert rows["U-004"]["platform"] is True and rows["T-001"]["platform"] is False
    assert rows["T-001"]["scope"] == "goods" and rows["U-001"]["scope"] == "bidder"

    cell_a = rows["U-001"]["cells"][str(_pkg(seeded, "A"))]
    assert cell_a["status"] == "meets" and cell_a["reason_short"] == ""
    assert cell_a["result_id"] is not None
    cell_c = rows["U-001"]["cells"][str(_pkg(seeded, "C"))]
    assert cell_c["status"] == "does_not_meet" and cell_c["reason_short"] == "Not provided by bidder"
    assert rows["U-001"]["cells"][str(_pkg(seeded, "D"))] == {"status": None, "reason_short": "", "result_id": None}
    assert rows["U-004"]["cells"][str(_pkg(seeded, "A"))]["reason_short"] == "Checked by the platform"


def test_filters_only_mandatory_scope_differences_and_problems(client: TestClient, seeded):
    def codes(query: str) -> list[str]:
        return [r["code"] for r in _comparison(client, seeded["project"], query)["rows"]]

    assert codes("?only_mandatory=true") == ["T-001", "U-001", "U-004"]
    assert codes("?scope=bidder") == ["U-001", "U-004", "U-002", "U-003"]
    assert codes("?scope=goods") == ["T-001"]
    # T-001 is the same for everyone who ran. U-004 is needs review for everyone, so it is not a difference.
    assert codes("?only_differences=true") == ["U-001", "U-002", "U-003"]
    # U-002 has partial and meets only: not a problem. U-004 is needs review for everyone: a problem.
    assert codes("?only_problems=true") == ["U-001", "U-004", "U-003"]
    assert codes("?only_mandatory=true&only_problems=true") == ["U-001", "U-004"]


def test_rows_are_grouped_by_scope_then_category(client: TestClient, seeded):
    body = _comparison(client, seeded["project"])
    categories = [(r["scope"], r["category"]) for r in body["rows"]]
    # Each category is one contiguous block.
    blocks = [key for index, key in enumerate(categories) if index == 0 or categories[index - 1] != key]
    assert len(blocks) == len(set(categories))
    assert blocks == [("goods", "Technical"), ("bidder", "Documents"), ("bidder", "Qualification"),
                      ("bidder", "Commercial")]


def test_filters_do_not_change_the_bidders(client: TestClient, seeded):
    full = _comparison(client, seeded["project"])
    filtered = _comparison(client, seeded["project"], "?only_problems=true&scope=bidder")
    assert full["bidders"] == filtered["bidders"]


# ---------- comparison: warnings ----------

def test_warning_for_a_bidder_without_a_finished_run(client: TestClient, seeded):
    warnings = _comparison(client, seeded["project"])["warnings"]
    assert [(w["code"], w["package_id"]) for w in warnings] == [("no_finished_run", _pkg(seeded, "D"))]


def test_warning_for_a_run_that_is_not_reproducible(client: TestClient, seeded):
    with session_scope() as s:
        run = s.get(Run, seeded["runs"]["C"])
        run.reproducible = False
        s.add(run)
    body = _comparison(client, seeded["project"])
    assert ("not_reproducible", _pkg(seeded, "C")) in [(w["code"], w["package_id"]) for w in body["warnings"]]
    assert _bidder(body, _pkg(seeded, "C"))["reproducible"] is False


def test_warning_when_bidders_were_checked_with_different_models(client: TestClient, seeded):
    with session_scope() as s:
        run = s.get(Run, seeded["runs"]["E"])
        run.model_config_used = {"match": "other-model"}
        s.add(run)
    codes = [w["code"] for w in _comparison(client, seeded["project"])["warnings"]]
    assert "model_config_mismatch" in codes


def test_no_config_warning_when_all_runs_agree(client: TestClient, seeded):
    codes = [w["code"] for w in _comparison(client, seeded["project"])["warnings"]]
    assert "model_config_mismatch" not in codes


def test_comparison_without_a_scoring_table_has_no_points(client: TestClient):
    ids = _seed(scheme=False)
    body = _comparison(client, ids["project"])
    assert all(b["points"] is None for b in body["bidders"])
    # Without points, the ranking uses the compliance lower bounds only: A 100, B 54.5, E 45.5.
    assert [b["package_id"] for b in body["bidders"][:3]] == [_pkg(ids, "A"), _pkg(ids, "B"), _pkg(ids, "E")]
    assert [b["rank"] for b in body["bidders"][:3]] == [1, 2, 3]


def test_russian_project_gets_russian_committee_wording(client: TestClient):
    ids = _seed(lang="ru")
    body = _comparison(client, ids["project"])
    assert body["advisory_label"] == "Рекомендательно: решение принимает комиссия"
    c = _bidder(body, _pkg(ids, "C"))
    assert c["verdict_committee_text"] == "Подлежит отклонению по решению комиссии"


# ---------- committee decisions ----------

def test_committee_decision_is_logged_and_listed(client: TestClient, seeded):
    before = _comparison(client, seeded["project"])
    url = f"{API}/projects/{seeded['project']}/bidders/{_pkg(seeded, 'B')}/decision"
    resp = client.post(url, json={"member_name": "Committee member 1", "text": "  Asked for the missing page.  "})
    assert resp.status_code == 201, resp.text
    row = resp.json()
    assert row["package_id"] == _pkg(seeded, "B")
    assert row["member_name"] == "Committee member 1"
    assert row["text"] == "Asked for the missing page."

    history = client.get(url).json()
    assert [h["text"] for h in history] == ["Asked for the missing page."]

    with session_scope() as s:
        logged = s.exec(select(AuditLog).where(AuditLog.action == "committee_decision")).all()
        assert len(logged) == 1 and logged[0].entity_id == _pkg(seeded, "B")

    assert _comparison(client, seeded["project"])["bidders"] == before["bidders"]  # nothing else changed


def test_committee_decision_validation_and_bidder_check(client: TestClient, seeded):
    url = f"{API}/projects/{seeded['project']}/bidders/{_pkg(seeded, 'B')}/decision"
    assert client.post(url, json={"member_name": "  ", "text": "x"}).status_code == 422
    assert client.post(url, json={"member_name": "Member", "text": "   "}).status_code == 422
    assert client.get(url).json() == []
    tender_url = f"{API}/projects/{seeded['project']}/bidders/{seeded['tender']}/decision"
    assert client.post(tender_url, json={"member_name": "Member", "text": "x"}).status_code == 404


# ---------- XLSX ----------

def test_xlsx_has_matrix_and_summary_with_the_label_and_no_quotes(client: TestClient, seeded):
    resp = client.get(f"{API}/projects/{seeded['project']}/comparison.xlsx")
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Matrix", "Summary"]

    for name in wb.sheetnames:
        assert wb[name]["A1"].value == ADVISORY_EN

    matrix = wb["Matrix"]
    header = [cell.value for cell in matrix[2]]
    assert header[:6] == ["Code", "Scope", "Type", "Category", "Requirement", "Verified by"]
    assert BIDDERS["A"] in header
    assert [matrix.cell(row=r, column=1).value for r in range(3, 8)] == ["T-001", "U-001", "U-004", "U-002", "U-003"]

    # Status colour from the web chips (does not meet is red); the reason is short text.
    col = header.index(BIDDERS["C"]) + 1
    cell = matrix.cell(row=4, column=col)  # U-001 row
    assert cell.value.startswith("Does not meet") and "Not provided by bidder" in cell.value
    assert cell.fill.fgColor.rgb.endswith("FDE4E1")

    summary = wb["Summary"]
    rows = [[c.value for c in row] for row in summary.iter_rows(min_row=3)]
    by_name = {row[1]: row for row in rows if row[1]}
    assert by_name[BIDDERS["A"]][0] == 1 and by_name[BIDDERS["A"]][2] == "123456789"
    assert by_name[BIDDERS["C"]][0] == "–"
    assert "U-001" in by_name[BIDDERS["C"]][16]  # deciding item codes

    everything = [str(c.value) for sheet in wb.worksheets for row in sheet.iter_rows() for c in row if c.value]
    assert not any("ZEBRA" in value for value in everything)  # no evidence quotes in the export


def test_xlsx_of_an_uzbek_project_is_in_uzbek(client: TestClient):
    ids = _seed(lang="uz")
    resp = client.get(f"{API}/projects/{ids['project']}/comparison.xlsx")
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content))
    matrix = wb["Matrix"]
    assert matrix["A1"].value == "Tavsiyaviy: qarorni baholash komissiyasi qabul qiladi"
    header = [cell.value for cell in matrix[2]]
    assert header[:6] == ["Kod", "Boʻlim", "Turi", "Toifa", "Talab", "Tekshiradi"]
    col = header.index(BIDDERS["C"]) + 1
    cell = matrix.cell(row=4, column=col).value  # U-001 row, bidder C
    assert cell.startswith("Mos kelmaydi") and "Ishtirokchi tomonidan taqdim etilmagan" in cell
    everything = " ".join(str(c.value) for sheet in wb.worksheets for row in sheet.iter_rows() for c in row if c.value)
    assert "Does not meet" not in everything and "Not provided" not in everything


def test_xlsx_without_an_approved_checklist_is_409(client: TestClient):
    ids = _bare_customer(approved=False)
    assert client.get(f"{API}/projects/{ids['project']}/comparison.xlsx").status_code == 409
