"""Clarification request as a Word letter (GET .../export/clarification.docx), synthetic data only.

Seeding reuses the results test rows (one tender with a scheme, one bid run). The questions added here are
hand-written: a contradiction with two codes and quotes, and a scheme question without a code.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlmodel import select

from app.db import session_scope
from app.models import Project, Question
from engine.letter_docx import DOCX_MIME
from tests.unit.test_results_api import _seed

API = "/api/v1"
TITLES = {
    "ru": "Запрос на разъяснение документации",
    "uz": "Tender hujjatiga tushuntirish soʻrovi",
    "en": "Request for clarification of the tender documentation",
}
Q_CONTRADICTION = "Max score 15 in the tender conflicts with the criteria sum; please confirm the figure."
Q_SCHEME = "The criteria maxima add up to 14, but the stated maximum is 15."


@pytest.fixture
def seeded(client: TestClient) -> dict[str, Any]:
    ids = _seed(client)
    with session_scope() as s:
        project = s.get(Project, ids["project"])
        project.customer = "Тестовый заказчик"
        project.lot_number = "12"
        s.add(Question(
            project_id=ids["project"], kind="contradiction", text=Q_CONTRADICTION,
            sources=[
                {"page_id": ids["tpage28"], "quote": "Maximum score 15", "code": "U-001"},
                {"page_id": ids["tpage28"], "quote": "Rated power not less than 2.5 kW", "code": "T-001"},
            ],
        ))
        s.add(Question(
            project_id=ids["project"], kind="scheme_mismatch", text=Q_SCHEME,
            sources=[{"page_id": ids["tpage28"], "quote": "Maximum score 15"}],
        ))
    return ids


def _docx(client: TestClient, project_id: int) -> tuple[Any, Document]:
    resp = client.get(f"{API}/projects/{project_id}/export/clarification.docx")
    assert resp.status_code == 200, resp.text
    return resp, Document(io.BytesIO(resp.content))


def _texts(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs]


def test_docx_endpoint_returns_a_word_attachment_that_python_docx_opens(client, seeded):
    resp, doc = _docx(client, seeded["project"])
    assert resp.headers["content-type"] == DOCX_MIME
    assert resp.headers["content-disposition"] == f'attachment; filename="clarification-{seeded["project"]}.docx"'
    assert resp.content[:2] == b"PK"  # a zip container, as every .docx is
    assert _texts(doc)[0] == TITLES["ru"]


def test_docx_letter_has_addressee_project_lot_date_intro_and_closing(client, seeded):
    _, doc = _docx(client, seeded["project"])
    texts = _texts(doc)
    assert "Кому: Тестовый заказчик" in texts
    assert "Проект: Lot 1" in texts
    assert "Лот № 12" in texts
    assert f"Дата: {date.today().strftime('%d.%m.%Y')}" in texts
    assert "Прошу разъяснить следующие положения документации:" in texts
    assert "С уважением," in texts
    assert any(t.startswith("_____") and "ФИО" in t for t in texts)


def test_docx_lists_every_question_with_code_text_and_source_in_italic_quote(client, seeded):
    _, doc = _docx(client, seeded["project"])
    texts = _texts(doc)

    contradiction = next(t for t in texts if Q_CONTRADICTION in t)
    assert contradiction.startswith("1.")
    assert "U-001, T-001" in contradiction
    scheme = next(t for t in texts if Q_SCHEME in t)
    assert scheme.startswith("2.")

    # Each source: "<file>, стр. 28" in plain text, then the quote in italics.
    source_lines = [t for t in texts if "tender.pdf, стр. 28" in t]
    assert len(source_lines) == 3
    italic_quotes = [
        run.text for p in doc.paragraphs for run in p.runs if run.italic
    ]
    assert "«Maximum score 15»" in italic_quotes
    assert "«Rated power not less than 2.5 kW»" in italic_quotes


def test_docx_question_without_code_is_written_without_one(client, seeded):
    _, doc = _docx(client, seeded["project"])
    scheme = next(p for p in doc.paragraphs if Q_SCHEME in p.text)
    assert scheme.runs[0].text.startswith("2.")
    assert not any(run.bold for run in scheme.runs)


def test_docx_titles_differ_by_report_language(client, seeded):
    titles = {}
    for lang in ("ru", "uz", "en"):
        with session_scope() as s:
            s.get(Project, seeded["project"]).report_lang = lang
        _, doc = _docx(client, seeded["project"])
        titles[lang] = _texts(doc)[0]
    assert titles == TITLES
    assert len(set(titles.values())) == 3


def test_docx_english_letter_uses_english_labels(client, seeded):
    with session_scope() as s:
        s.get(Project, seeded["project"]).report_lang = "en"
    _, doc = _docx(client, seeded["project"])
    texts = _texts(doc)
    assert "To: Тестовый заказчик" in texts
    assert "Lot No. 12" in texts
    assert "Sincerely," in texts
    assert any("file tender.pdf, p. 28" in t for t in texts)


def test_docx_without_customer_is_addressed_to_the_customer_generically(client, seeded):
    with session_scope() as s:
        s.get(Project, seeded["project"]).customer = None
    _, doc = _docx(client, seeded["project"])
    assert "Заказчику" in _texts(doc)


def test_docx_with_no_questions_says_so(client):
    with session_scope() as s:
        project = Project(name="Empty lot", report_lang="en")
        s.add(project)
        s.flush()
        project_id = project.id
    _, doc = _docx(client, project_id)
    texts = _texts(doc)
    assert TITLES["en"] == texts[0]
    assert "No clarification questions were identified for this project." in texts
    assert not any(t.startswith("1.\t") for t in texts)


def test_docx_endpoint_unknown_project_is_404(client):
    assert client.get(f"{API}/projects/999/export/clarification.docx").status_code == 404


def test_txt_endpoint_still_works_with_the_same_questions(client, seeded):
    resp = client.get(f"{API}/projects/{seeded['project']}/export/clarification.txt")
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("text/plain")
    assert resp.headers["content-disposition"] == f'attachment; filename="clarification-project-{seeded["project"]}.txt"'
    text = resp.text
    assert text.startswith(TITLES["ru"])
    assert f"1. {Q_CONTRADICTION}" in text
    assert f"2. {Q_SCHEME}" in text
    assert "«Maximum score 15»" in text


def test_summary_page_links_the_docx_letter(client, seeded):
    text = client.get(f"/runs/{seeded['run']}/summary").text
    assert f"/api/v1/projects/{seeded['project']}/export/clarification.docx" in text
    assert f"/api/v1/projects/{seeded['project']}/export/clarification.txt" not in text
