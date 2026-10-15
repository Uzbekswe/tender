"""Customer-side API: bidders, check-all, comparison, XLSX export, committee decisions.

See docs/customer-side-decisions.md §4-7. The routes work for any project; `mode` is stored and
shown, but nothing here refuses a supplier project. Everything the system outputs is advisory:
no route awards, rejects or changes a verdict, and a committee decision only writes an AuditLog row.
"""

from __future__ import annotations

import io
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlmodel import Session, col, select

from app.api.projects import start_run
from app.api.results import _clean
from app.db import get_session
from app.models import (
    AuditLog,
    Document,
    Package,
    PackageRole,
    Page,
    PageCard,
    Project,
    Run,
    RunStatus,
    Scope,
)
from app.schemas import PackageRead, RunCreate
from app.schemas_customer import (
    BidderCreate,
    BidderUpdate,
    CheckAllOut,
    CheckAllRun,
    CheckAllSkipped,
    ComparisonOut,
    DecisionIn,
    DecisionOut,
)
from app.storage import FileStore, get_store
from app.web.i18n import translator
from engine.comparison import ComparisonError, approved_version, bid_packages, build_comparison, has_files

router = APIRouter(prefix="/api/v1", tags=["customer"])

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_FINISHED = frozenset({RunStatus.done, RunStatus.failed, RunStatus.waiting_approval})
_IN_PROGRESS = frozenset({RunStatus.queued, RunStatus.running})
DECISION_ACTION = "committee_decision"

# Status chip colours, the same as app/web/static/app.css (background, text).
_STATUS_COLOURS = {
    "meets": ("E3F4E9", "1F7A3E"),
    "partial": ("FFF4CC", "8A6500"),
    "not_confirmed": ("FFE8D1", "B35400"),
    "does_not_meet": ("FDE4E1", "B42318"),
    "needs_review": ("E7EAEF", "515B68"),
}


def _val(x: Any) -> Any:
    return getattr(x, "value", x)


def _project_or_404(session: Session, project_id: int) -> Project:
    project = session.get(Project, project_id)
    if project is None:
        raise HTTPException(status_code=404, detail=f"project {project_id} not found")
    return project


def _bidder_or_404(session: Session, project_id: int, package_id: int) -> Package:
    package = session.get(Package, package_id)
    if package is None or package.project_id != project_id or _val(package.role) != PackageRole.bid.value:
        raise HTTPException(status_code=404, detail=f"bidder {package_id} not found in project {project_id}")
    return package


# ---------- bidders ----------

@router.post("/projects/{project_id}/bidders", response_model=PackageRead, status_code=201)
def create_bidder(project_id: int, body: BidderCreate, session: Session = Depends(get_session)) -> PackageRead:
    _project_or_404(session, project_id)
    package = Package(
        project_id=project_id,
        role=PackageRole.bid,
        label=body.name[:120],
        bidder_name=body.name,
        bidder_stir=body.stir,
    )
    session.add(package)
    session.flush()
    session.add(AuditLog(
        action="bidder.create", entity="package", entity_id=package.id,
        before={}, after={"name": body.name, "stir": body.stir},
    ))
    session.commit()
    session.refresh(package)
    return PackageRead.model_validate(package)


@router.get("/projects/{project_id}/bidders", response_model=list[PackageRead])
def list_bidders(project_id: int, session: Session = Depends(get_session)) -> list[PackageRead]:
    _project_or_404(session, project_id)
    return [PackageRead.model_validate(p) for p in bid_packages(session, project_id)]


@router.patch("/projects/{project_id}/bidders/{package_id}", response_model=PackageRead)
def update_bidder(
    project_id: int, package_id: int, body: BidderUpdate, session: Session = Depends(get_session),
) -> PackageRead:
    package = _bidder_or_404(session, project_id, package_id)
    changes = body.model_dump(exclude_unset=True)
    if "name" in changes and changes["name"] is None:
        raise HTTPException(status_code=422, detail="name cannot be null")
    before = {"name": package.bidder_name, "stir": package.bidder_stir}
    if "name" in changes:
        package.bidder_name = changes["name"]
        package.label = changes["name"][:120]
    if "stir" in changes:
        package.bidder_stir = changes["stir"]
    after = {"name": package.bidder_name, "stir": package.bidder_stir}
    if after != before:
        session.add(package)
        session.add(AuditLog(
            action="bidder.update", entity="package", entity_id=package.id, before=before, after=after,
        ))
        session.commit()
        session.refresh(package)
    return PackageRead.model_validate(package)


@router.delete("/projects/{project_id}/bidders/{package_id}", status_code=204)
def delete_bidder(
    project_id: int,
    package_id: int,
    session: Session = Depends(get_session),
    store: FileStore = Depends(get_store),
) -> Response:
    """Only a bidder with no check at all can be deleted. Its uploads, pages and page cards go with it.

    A finished check (done, failed or waiting for approval) blocks the delete (409), and so does a
    check still queued or running. The audit log keeps the deletion.
    """
    package = _bidder_or_404(session, project_id, package_id)
    statuses = [RunStatus(s) for s in session.exec(select(Run.status).where(Run.package_id == package.id)).all()]
    if any(s in _FINISHED for s in statuses):
        raise HTTPException(status_code=409, detail="bidder has a finished check and cannot be deleted")
    if any(s in _IN_PROGRESS for s in statuses):
        raise HTTPException(status_code=409, detail="bidder has a check in progress")

    documents = list(session.exec(select(Document).where(Document.package_id == package.id)).all())
    doc_ids = [d.id for d in documents]
    pages = list(session.exec(select(Page).where(col(Page.document_id).in_(doc_ids))).all()) if doc_ids else []
    page_ids = [p.id for p in pages]
    for page in pages:
        page.duplicate_of_page_id = None  # duplicates may point at each other inside the package
        session.add(page)
    session.flush()
    if page_ids:
        for card in session.exec(select(PageCard).where(col(PageCard.page_id).in_(page_ids))).all():
            session.delete(card)
    session.flush()
    for page in pages:
        session.delete(page)
    session.flush()
    for document in documents:
        session.delete(document)
    session.flush()

    files = [d.stored_path for d in documents] + [d.pdf_path for d in documents] + [p.image_path for p in pages]
    session.add(AuditLog(
        action="bidder.delete", entity="package", entity_id=package.id,
        before={"name": package.bidder_name, "stir": package.bidder_stir, "documents": len(documents)},
        after={},
    ))
    session.delete(package)
    session.commit()

    for rel_path in files:
        if rel_path:
            try:
                store.delete(rel_path)
            except ValueError:  # outside the storage root: nothing of ours to remove
                pass
    return Response(status_code=204)


# ---------- check all ----------

@router.post("/projects/{project_id}/check-all", response_model=CheckAllOut, status_code=202)
def check_all(project_id: int, session: Session = Depends(get_session)) -> CheckAllOut:
    """One bid run per bidder with files, all on the latest approved checklist version.

    Reuses the run creation of POST /packages/{id}/runs. A bidder whose check on this version is still
    queued or running is skipped, so a double click does not queue the same work twice. Later approved
    versions re-run everyone, because the latest approved version is always used.
    """
    _project_or_404(session, project_id)
    version = approved_version(session, project_id)
    if version is None:
        raise HTTPException(
            status_code=409, detail="no approved checklist for this project; approve the tender checklist first",
        )
    bidders = bid_packages(session, project_id)
    if not bidders:
        raise HTTPException(status_code=409, detail="the project has no bidders")

    runs: list[CheckAllRun] = []
    skipped: list[CheckAllSkipped] = []
    for package in bidders:
        if not has_files(session, package.id):
            skipped.append(CheckAllSkipped(package_id=package.id, reason="no files uploaded"))
            continue
        active = session.exec(
            select(Run.id).where(
                Run.package_id == package.id,
                Run.checklist_version_id == version.id,
                col(Run.status).in_(list(_IN_PROGRESS)),
            ).limit(1)
        ).first()
        if active is not None:
            skipped.append(CheckAllSkipped(package_id=package.id, reason="a check of this bidder is already running"))
            continue
        accepted = start_run(package.id, RunCreate(checklist_version_id=version.id), session)
        runs.append(CheckAllRun(package_id=package.id, run_id=accepted.run_id))

    return CheckAllOut(
        checklist_version_id=version.id,
        checklist_version_number=version.number,
        runs=runs,
        skipped=skipped,
    )


# ---------- comparison ----------

def _comparison(
    session: Session, project_id: int, *, only_mandatory: bool, only_differences: bool,
    only_problems: bool, scope: Scope | None,
) -> tuple[dict[str, Any], Any]:
    project = _project_or_404(session, project_id)
    t = translator(project.report_lang or "ru")
    try:
        data = build_comparison(
            session, project_id, t,
            only_mandatory=only_mandatory,
            only_differences=only_differences,
            only_problems=only_problems,
            scope=None if scope is None else _val(scope),
        )
    except ComparisonError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    return data, t


@router.get("/projects/{project_id}/comparison", response_model=ComparisonOut)
def get_comparison(
    project_id: int,
    only_mandatory: bool = False,
    only_differences: bool = False,
    only_problems: bool = False,
    scope: Scope | None = None,
    session: Session = Depends(get_session),
) -> ComparisonOut:
    data, _ = _comparison(
        session, project_id, only_mandatory=only_mandatory, only_differences=only_differences,
        only_problems=only_problems, scope=scope,
    )
    return ComparisonOut.model_validate(data)


# ---------- XLSX ----------

def _cell_text(cell: dict[str, Any], t: Any) -> str:
    if cell.get("status") is None:
        return t("committee.xlsx.none")
    label = t(f"status.{cell['status']}")
    reason = cell.get("reason_short") or ""
    return f"{label}\n{reason}" if reason else label


def _style_header(ws: Any, row: int) -> None:
    for cell in ws[row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _write_table(ws: Any, label: str, header: list[str], body: list[list[Any]], widths: list[int]) -> int:
    """Row 1: the advisory label. Row 2: the header. Then the body. Returns the next free row."""
    ws.append([_clean(label)])
    ws["A1"].font = Font(bold=True, italic=True)
    ws.append([_clean(h) for h in header])
    _style_header(ws, 2)
    for row in body:
        ws.append([_clean(v) if isinstance(v, str) else v for v in row])
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A3"
    return ws.max_row + 1


def build_comparison_workbook(data: dict[str, Any], t: Any) -> bytes:
    """Sheets "Matrix" (requirements x bidders, status and short reason) and "Summary" (one row per bidder).

    Neither sheet has personal-name columns or evidence quotes (docs/customer-side-decisions.md §7.1, risk 2).
    The matrix uses the same status colours as the web chips.
    """
    label = data["advisory_label"]
    bidders = data["bidders"]
    wb = Workbook()

    matrix = wb.active
    matrix.title = "Matrix"
    header = [
        t("committee.xlsx.code"), t("committee.xlsx.scope"), t("committee.xlsx.type"),
        t("committee.xlsx.category"), t("committee.xlsx.parameter"), t("committee.xlsx.verified_by"),
    ] + [b["name"] for b in bidders]
    body: list[list[Any]] = []
    fills: list[tuple[int, int, str]] = []  # (excel row, excel column, status)
    for row in data["rows"]:
        values: list[Any] = [
            row["code"], t(f"checklist.scope.{row['scope']}"), t(f"req_type.{row['req_type']}"),
            row["category"], row["name"],
            t("checklist.verification.platform" if row["platform"] else "checklist.verification.documents"),
        ]
        for col_index, bidder in enumerate(bidders, start=len(values) + 1):
            cell = row["cells"].get(bidder["package_id"], {})
            values.append(_cell_text(cell, t))
            if cell.get("status"):
                fills.append((len(body) + 3, col_index, cell["status"]))
        body.append(values)
    next_row = _write_table(matrix, label, header, body, [10, 16, 14, 22, 40, 16] + [30] * len(bidders))
    for excel_row, excel_col, status in fills:
        background, colour = _STATUS_COLOURS.get(status, ("FFFFFF", "1D232B"))
        target = matrix.cell(row=excel_row, column=excel_col)
        target.fill = PatternFill("solid", fgColor=background)
        target.font = Font(color=colour)
    if not data["rows"]:
        matrix.cell(row=next_row, column=1, value=t("committee.xlsx.none"))

    summary = wb.create_sheet("Summary")
    summary_header = [
        t("committee.xlsx.rank"), t("committee.xlsx.bidder"), t("committee.xlsx.stir"),
        t("committee.xlsx.verdict"), t("committee.xlsx.points_min"), t("committee.xlsx.points_max"),
        t("committee.xlsx.pass_mark"), t("committee.xlsx.compliance_min"), t("committee.xlsx.compliance_max"),
        t("committee.xlsx.compliance"), t("committee.xlsx.mandatory_met"),
        t("status.meets"), t("status.partial"), t("status.not_confirmed"), t("status.does_not_meet"),
        t("status.needs_review"), t("committee.xlsx.deciding"), t("committee.xlsx.order"),
        t("committee.xlsx.reproducible"),
    ]
    summary_body: list[list[Any]] = []
    for bidder in bidders:
        comp = bidder["compliance"] or {}
        points = bidder["points"] or {}
        counts = bidder["counts"] or {}
        mandatory = (
            f"{bidder['mandatory_met']} of {bidder['mandatory_total']}"
            if bidder["mandatory_total"] is not None else ""
        )
        summary_body.append([
            bidder["rank"] if bidder["rank"] is not None else "–",
            bidder["name"],
            bidder["stir"] or "",
            bidder["verdict_committee_text"] or "",
            points.get("min", ""), points.get("max", ""), points.get("pass_mark") or "",
            comp.get("min", ""), comp.get("max", ""), comp.get("point", ""),
            mandatory,
            counts.get("meets", ""), counts.get("partial", ""), counts.get("not_confirmed", ""),
            counts.get("not_met", ""), counts.get("needs_review", ""),
            ", ".join(item["code"] for item in bidder["deciding_items"]),
            t("committee.order_not_certain") if bidder["order_not_certain"] else "",
            "" if bidder["reproducible"] is None else t("common.yes" if bidder["reproducible"] else "common.no"),
        ])
    next_row = _write_table(
        summary, label, summary_header, summary_body,
        [8, 32, 12, 40, 12, 12, 12, 14, 14, 14, 16, 10, 10, 14, 14, 14, 24, 20, 14],
    )
    if data["warnings"]:
        next_row += 1
        summary.cell(row=next_row, column=1, value=t("committee.xlsx.warnings")).font = Font(bold=True)
        for warning in data["warnings"]:
            next_row += 1
            summary.cell(row=next_row, column=1, value=_clean(warning["text"]))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@router.get("/projects/{project_id}/comparison.xlsx")
def export_comparison(
    project_id: int,
    only_mandatory: bool = False,
    only_differences: bool = False,
    only_problems: bool = False,
    scope: Scope | None = None,
    session: Session = Depends(get_session),
) -> Response:
    data, t = _comparison(
        session, project_id, only_mandatory=only_mandatory, only_differences=only_differences,
        only_problems=only_problems, scope=scope,
    )
    body = build_comparison_workbook(data, t)
    return Response(
        content=body,
        media_type=XLSX_MIME,
        headers={"Content-Disposition": f'attachment; filename="comparison-project-{project_id}.xlsx"'},
    )


# ---------- committee decisions ----------

def _decision_out(row: AuditLog) -> DecisionOut:
    after = dict(row.after or {})
    return DecisionOut(
        id=row.id,
        package_id=row.entity_id or 0,
        member_name=str(after.get("member_name") or row.who),
        text=str(after.get("text") or ""),
        at=row.at,
    )


@router.post("/projects/{project_id}/bidders/{package_id}/decision", response_model=DecisionOut, status_code=201)
def add_decision(
    project_id: int, package_id: int, body: DecisionIn, session: Session = Depends(get_session),
) -> DecisionOut:
    """Free-text committee decision on one bidder. Logged in the audit log; nothing else changes."""
    package = _bidder_or_404(session, project_id, package_id)
    row = AuditLog(
        who=body.member_name,
        action=DECISION_ACTION,
        entity="package",
        entity_id=package.id,
        before={},
        after={"member_name": body.member_name, "text": body.text},
    )
    session.add(row)
    session.commit()
    session.refresh(row)
    return _decision_out(row)


@router.get("/projects/{project_id}/bidders/{package_id}/decision", response_model=list[DecisionOut])
def list_decisions(
    project_id: int, package_id: int, session: Session = Depends(get_session),
) -> list[DecisionOut]:
    package = _bidder_or_404(session, project_id, package_id)
    rows = session.exec(
        select(AuditLog)
        .where(AuditLog.action == DECISION_ACTION, AuditLog.entity == "package", AuditLog.entity_id == package.id)
        .order_by(AuditLog.id)
    ).all()
    return [_decision_out(row) for row in rows]
